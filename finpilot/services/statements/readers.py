"""Turn any statement file into grids of cells: CSV/TSV/TXT, Excel (.xlsx/.xls), HTML tables, PDF.

Format is detected from the file's bytes first (banks routinely mislabel files — many Indian banks
serve an HTML table named `.xls`), and from the extension only as a tie-breaker.
"""

from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass, field
from html.parser import HTMLParser

from finpilot.services.statements.values import Cell

log = logging.getLogger(__name__)

Grid = list[list[Cell]]
SUPPORTED_EXTENSIONS = (".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".xls", ".htm", ".html", ".pdf")


class StatementError(ValueError):
    """The file can't be read at all (as opposed to individual rows being invalid)."""

    def __init__(self, message: str, code: str = "unreadable"):
        super().__init__(message)
        self.code = code


@dataclass
class Table:
    grid: Grid
    source: str  # e.g. "csv", "xlsx:Sheet1", "pdf:page 2"


@dataclass
class Document:
    format: str
    tables: list[Table] = field(default_factory=list)
    text_lines: list[str] = field(default_factory=list)  # PDF text fallback (no ruled tables)


def detect_format(content: bytes, filename: str | None = None) -> str:
    head = content[:2048].lstrip()
    if content[:5] == b"%PDF-":
        return "pdf"
    if content[:4] == b"PK\x03\x04":
        return "xlsx"
    if content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    low = head.lower()
    if low.startswith((b"<!doctype html", b"<html", b"<table", b"<?xml")) or b"<table" in low[:2048]:
        return "html"
    ext = (filename or "").lower().rsplit(".", 1)[-1] if filename and "." in filename else ""
    if ext == "pdf":
        raise StatementError(
            "The file is named .pdf but is not a valid PDF. Re-download the statement from your bank.",
            code="corrupt",
        )
    if b"\x00" in content[:4096] and ext not in ("csv", "txt", "tsv"):
        raise StatementError(
            f"Unrecognised binary file{f' (.{ext})' if ext else ''}. Supported: CSV, TSV, TXT, XLSX, XLS, HTML, PDF.",
            code="unsupported",
        )
    return "csv"  # includes CSV files that banks label .xls


def read(content: bytes, filename: str | None = None, password: str | None = None) -> Document:
    fmt = detect_format(content, filename)
    if fmt == "pdf":
        return _read_pdf(content, password)
    if fmt == "xlsx":
        return Document("xlsx", _read_xlsx(content))
    if fmt == "xls":
        return Document("xls", _read_xls(content))
    if fmt == "html":
        return Document("html", _read_html(_decode(content)))
    return Document("csv", [Table(_read_csv(_decode(content)), "csv")])


# ----------------------------------------------------------------------------- text formats
def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = content.decode(enc)
        except UnicodeDecodeError:
            continue
        if enc == "utf-16" and not content.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        return text
    return content.decode("utf-8", errors="replace")


def _read_csv(text: str) -> Grid:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    sample = text[:16384]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in (",", ";", "\t", "|")}
        delimiter = max(counts, key=lambda d: counts[d]) if any(counts.values()) else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


class _TableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[Grid] = []
        self._stack: list[Grid] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._stack.append([])
        elif tag == "tr" and self._stack:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._stack and self._row is not None:
            if any(c for c in self._row):
                self._stack[-1].append(self._row)
            self._row = None
        elif tag == "table" and self._stack:
            grid = self._stack.pop()
            if grid:
                self.tables.append(grid)

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _read_html(text: str) -> list[Table]:
    parser = _TableParser()
    parser.feed(text)
    tables = [Table(g, f"html:table {i + 1}") for i, g in enumerate(parser.tables)]
    if not tables:
        raise StatementError("No table found in the HTML/XLS file.", code="no_table")
    return tables


# ----------------------------------------------------------------------------- Excel
def _read_xlsx(content: bytes) -> list[Table]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency is in requirements.txt
        raise StatementError("Excel support needs the 'openpyxl' package.", code="dependency") from exc
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - zip/xml errors from a corrupt file
        raise StatementError(
            f"Could not open the Excel file ({type(exc).__name__}).", code="corrupt"
        ) from exc
    tables = []
    for ws in wb.worksheets:
        grid = [list(row) for row in ws.iter_rows(values_only=True)]
        if any(any(c is not None and str(c).strip() for c in r) for r in grid):
            tables.append(Table(grid, f"xlsx:{ws.title}"))
    wb.close()
    return tables


def _read_xls(content: bytes) -> list[Table]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover
        raise StatementError("Legacy .xls support needs the 'xlrd' package.", code="dependency") from exc
    try:
        book = xlrd.open_workbook(file_contents=content)
    except Exception as exc:  # noqa: BLE001
        raise StatementError(f"Could not open the .xls file ({type(exc).__name__}).", code="corrupt") from exc
    tables = []
    for sheet in book.sheets():
        grid: Grid = []
        for r in range(sheet.nrows):
            row: list[Cell] = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    except Exception:  # noqa: BLE001
                        row.append(cell.value)
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    row.append(None)
                else:
                    row.append(cell.value)
            grid.append(row)
        if grid:
            tables.append(Table(grid, f"xls:{sheet.name}"))
    return tables


# ----------------------------------------------------------------------------- PDF
def _read_pdf(content: bytes, password: str | None) -> Document:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover
        raise StatementError("PDF support needs the 'pdfplumber' package.", code="dependency") from exc
    doc = Document("pdf")
    try:
        pdf = pdfplumber.open(io.BytesIO(content), password=password or "")
        _ = len(pdf.pages)  # some encrypted files only fail when pages are touched
    except Exception as exc:  # noqa: BLE001 - pdfplumber wraps pdfminer errors (PdfminerException(PDFPasswordIncorrect()))
        if "PasswordIncorrect" in repr(exc) or "password" in str(exc).lower():
            raise StatementError(
                "This PDF is password-protected. Enter the statement password (banks often use your date of "
                "birth or customer ID)."
                if not password
                else "That password did not open the PDF.",
                code="password_required" if not password else "password_incorrect",
            ) from exc
        raise StatementError(f"Could not open the PDF ({type(exc).__name__}).", code="corrupt") from exc
    with pdf:
        for no, page in enumerate(pdf.pages, start=1):
            found = False
            for settings in (
                {"vertical_strategy": "lines", "horizontal_strategy": "lines"},
                {"vertical_strategy": "text", "horizontal_strategy": "text", "snap_tolerance": 3},
            ):
                try:
                    grids = page.extract_tables(table_settings=settings)
                except Exception:  # noqa: BLE001 - malformed page objects
                    grids = []
                grids = [g for g in grids if g and len(g) >= 2 and max(len(r) for r in g) >= 3]
                if grids:
                    for g in grids:
                        doc.tables.append(Table([list(r) for r in g], f"pdf:page {no}"))
                    found = True
                    break
            text = page.extract_text() or ""
            doc.text_lines += [ln.strip() for ln in text.splitlines() if ln.strip()]
            if not found:
                log.debug("pdf page %s: no table detected", no)
    if not doc.tables and not doc.text_lines:
        raise StatementError(
            "No text found in the PDF — it is probably a scanned image. Download the e-statement (not a "
            "scan) or export CSV/Excel from net banking.",
            code="scanned_pdf",
        )
    return doc


def merge_tables(tables: list[Table]) -> list[Table]:
    """PDF statements repeat one table per page: stitch same-width tables into one grid."""
    if len(tables) <= 1:
        return tables
    width = lambda t: max((len(r) for r in t.grid), default=0)  # noqa: E731
    merged: list[Table] = []
    for t in tables:
        if (
            merged
            and width(merged[-1]) == width(t)
            and merged[-1].source.split(":")[0] == t.source.split(":")[0]
        ):
            merged[-1] = Table(merged[-1].grid + t.grid, merged[-1].source + "+")
        else:
            merged.append(Table(list(t.grid), t.source))
    return merged
